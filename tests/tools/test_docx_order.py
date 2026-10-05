# Copyright 2023 The Qwen team, Alibaba Group. All rights reserved.
#
# Licensed under the Apache License, Version 2.0 (the "License");
# you may not use this file except in compliance with the License.
# You may obtain a copy of the License at
#
# http://www.apache.org/licenses/LICENSE-2.0
#
# Unless required by applicable law or agreed to in writing, software
# distributed under the License is distributed on an "AS IS" BASIS,
# WITHOUT WARRANTIES OR CONDITIONS OF ANY KIND, either express or implied.
# See the License for the specific language governing permissions and
# limitations under the License.

import pytest
from docx import Document

from qwen_agent.tools.simple_doc_parser import SimpleDocParser, parse_word


@pytest.fixture
def interleaved_docx(tmp_path):
    document = Document()
    document.add_paragraph('表格前\n第一行')
    first = document.add_table(rows=1, cols=2)
    first.cell(0, 0).text = '甲\n乙'
    first.cell(0, 1).text = '丙'
    document.add_paragraph('')
    document.add_paragraph('两个表格之间')
    second = document.add_table(rows=1, cols=2)
    second.cell(0, 0).text = '丁'
    second.cell(0, 1).text = '戊'
    document.add_paragraph('表格后')
    path = tmp_path / 'interleaved.docx'
    document.save(path)
    expected = [{
        'text': '表格前\n第一行'
    }, {
        'table': '|甲\n乙|丙|'
    }, {
        'text': ''
    }, {
        'text': '两个表格之间'
    }, {
        'table': '|丁|戊|'
    }, {
        'text': '表格后'
    }]
    return path, expected


def test_interleaved_blocks_preserve_order(interleaved_docx):
    path, expected = interleaved_docx
    assert parse_word(str(path)) == [{'page_num': 1, 'content': expected}]


@pytest.mark.parametrize('kind', ['paragraphs', 'tables', 'empty'])
def test_single_block_kind(tmp_path, kind):
    document = Document()
    expected = []
    if kind == 'paragraphs':
        for text in ['first', '', 'last']:
            document.add_paragraph(text)
            expected.append({'text': text})
    elif kind == 'tables':
        for text in ['first', 'last']:
            document.add_table(rows=1, cols=1).cell(0, 0).text = text
            expected.append({'table': f'|{text}|'})
    path = tmp_path / f'{kind}.docx'
    document.save(path)
    assert parse_word(str(path)) == [{'page_num': 1, 'content': expected}]


@pytest.mark.parametrize('structured_doc', [True, False])
def test_sdk_and_cache_preserve_order(tmp_path, interleaved_docx, structured_doc):
    path, expected = interleaved_docx
    parser = SimpleDocParser({'path': str(tmp_path / 'cache'), 'structured_doc': structured_doc})
    first = parser.call({'url': str(path)})
    if structured_doc:
        assert [{
            'text': block['text']
        } if 'text' in block else {
            'table': block['table']
        } for block in first[0]['content']] == expected
    else:
        # Empty paragraphs remain in the plain output, between the same blocks.
        assert first == '\n'.join(block.get('text', block.get('table')) for block in expected)
    # A cached parse must retain the same order even when the file is absent.
    path.unlink()
    assert parser.call({'url': str(path)}) == first


def test_extract_image_remains_unsupported(tmp_path):
    with pytest.raises(ValueError, match='Currently, extracting images is not supported!'):
        parse_word(str(tmp_path / 'missing.docx'), extract_image=True)
